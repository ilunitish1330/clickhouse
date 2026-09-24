#!/usr/bin/env python3
"""Load ANY raw data into ClickHouse and aggregate it -- with nobody describing it.

    python3 auto_load.py billing --file "Statistic Report.zip"     # .csv .tsv .txt .xlsx or a .zip of them
    python3 auto_load.py billing --sql-file query.sql              # a query run on SQL Server (.env)
    python3 auto_load.py billing --sql "SELECT * FROM X" --db PUC  # inline, another database
    python3 auto_load.py --inbox ~/inbox                           # every new file dropped in a folder
    python3 auto_load.py billing --rebuild                         # rebuild from what is loaded
    python3 auto_load.py billing --remodel                         # forget the model, decide again

What happens, in ClickHouse database <name>:
  1. raw        every row exactly as received, all text. A file replaces an
                earlier file of the same name; a query replaces everything.
  2. model      the first load profiles every column and DECIDES: measures,
                dates, dimensions (code/description pairs and hierarchies found
                from the data), entity keys, document numbers, junk to drop, and
                the aggregate grains. Saved to models/<name>.json and explained in
                models/<name>.md with a list of risks. Later loads REUSE it, so the
                structure never shifts under a report. Edit the json + --rebuild
                to overrule a decision.
  3. fact, dim_*, agg_*   rebuilt from raw by the model, every load.

Inbox: files in <inbox>/<name>/ load into dataset <name>; files directly in
<inbox> get a name from the file name, digits dropped ("Statistic Report
24092026.csv" -> statistic_report). Loaded files move to <inbox>/done/, failures
to <inbox>/failed/ with the error. Run it from cron every 15 minutes.

No human input is needed, but some things cannot be known from data alone
(what a code means, whether a credit should count). Those are not guessed:
they are listed under "Check" in the report.
"""
import argparse
import csv
import fcntl
import io
import json
import re
import shutil
import sys
import time
import zipfile
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import ax_load  # ch(), insert(), connect(), .env handling
from ax_load import ch

HERE = Path(__file__).resolve().parent
MODELS = HERE / "models"
CHUNK = 20_000
LOWCARD = 1000            # "few values" (a dimension attribute) caps here...
LOWCARD_SHARE = 0.05      # ...and at 5% of the rows, so 1,000 invoices in 3,000 rows are not "few"
FD_TOLERANCE = 0.001      # "A decides B" may be broken by 0.1% of A's values (data noise)
RESERVED = {"period", "_source", "_row", "_loaded_at", "row_count"}
PROTECTED = {"system", "default", "information_schema", "ax", "ub", "learn"}

# Name hints, checked on the upper-cased column name.
ID_SUFFIX = re.compile(r"(ID|RECID|KEY|NUM|NUMBER|NO|CODE|TYPE|STATUS|GROUP|CATEGORY|CATEGORIES|"
                       r"REGION|COUNTER|YEAR|MONTH|PERIOD|INDICATION|FLAG|ORIGIN)$")
MEASURE = re.compile(r"AMOUNT|AMT|QTY|QUANTITY|VALUE|COST|PRICE|TOTAL|SUM|BALANCE|TAX|DISC|"
                     r"REVENUE|SALES|WEIGHT|VOLUME|HOURS|KWH|CONSUMPTION")
NON_ADDITIVE = re.compile(r"PRICE|RATE|PERCENT|PCT|RATIO|AVG|AVERAGE|FACTOR")
QTY = re.compile(r"QTY|QUANTITY|VOLUME|WEIGHT|HOURS|KWH|UNITS|CONSUMPTION")
MONEY = re.compile(r"AMOUNT|AMT|VALUE|COST|TOTAL|REVENUE|SALES|TAX|BALANCE")
DOCUMENT = re.compile(r"INVOICE|VOUCHER|DOC|ORDER|TRANS|JOURNAL|LINE|RECEIPT|SLIP|REF|RECID")
DATEISH = re.compile(r"DATE|TIME|DAY|TICKS")
PREFERRED_DATE = re.compile(r"(INVOICE|TRANS|POSTING|DOC|ORDER|ACCOUNTING|BILL)\w*DATE")

DATE_RE = r"^(\\d{4}-\\d{1,2}-\\d{1,2}|\\d{1,2}[/.-]\\d{1,2}[/.-]\\d{2,4})"


# =============================================================================
# Reading the input
# =============================================================================

def cell(v) -> str:
    """Any source value as clean text. Dates keep full precision, 'NULL' is empty."""
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == datetime.min.time() else v.isoformat(sep=" ")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, bool):
        return "1" if v else "0"
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    if isinstance(v, (bytes, bytearray)):
        return v.hex()
    s = str(v).replace("\xa0", " ").strip()
    return "" if s.upper() == "NULL" else s


def decode(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")  # what Excel writes


def read_delimited(data: bytes):
    text = decode(data)
    first = text.split("\n", 1)[0]
    delim = max(["\t", ",", ";", "|"], key=first.count)
    reader = csv.reader(io.StringIO(text), delimiter=delim)
    return next(reader), ([cell(v) for v in r] for r in reader)


def read_xlsx(data: bytes):
    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.exit("reading .xlsx needs openpyxl:  pip install openpyxl")
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    # the biggest sheet is the data; small ones are usually notes or pivots
    ws = max(wb.worksheets, key=lambda s: s.max_row or 0)
    rows = ws.iter_rows(values_only=True)
    for header in rows:  # first row with at least two filled cells
        if sum(v is not None and str(v).strip() != "" for v in header) >= 2:
            return [cell(v) for v in header], ([cell(v) for v in r] for r in rows)
    return [], iter(())


def read_any(path: Path):
    """Yield (source name, header, rows) for a file, or for each file in a zip."""
    def one(name, data):
        if name.lower().endswith((".xlsx", ".xlsm")):
            return read_xlsx(data)
        if name.lower().endswith(".xls"):
            sys.exit(f"{name}: old .xls is not supported, save it as .xlsx or .csv")
        return read_delimited(data)

    if path.suffix.lower() == ".zip":
        with zipfile.ZipFile(path) as z:
            for name in z.namelist():
                if name.lower().endswith((".csv", ".tsv", ".txt", ".xlsx", ".xlsm", ".xls")):
                    yield (Path(name).name, *one(name, z.read(name)))
    else:
        yield (path.name, *one(path.name, path.read_bytes()))


def sanitize(names: list[str]) -> list[str]:
    """Column names Power BI and SQL are happy with: lower_snake, unique."""
    out, seen = [], set()
    for i, n in enumerate(names):
        s = re.sub(r"[^0-9a-zA-Z]+", "_", str(n)).strip("_").lower() or f"col_{i + 1}"
        if s[0].isdigit():
            s = "c_" + s
        if s in RESERVED:
            s += "_src"
        base, k = s, 2
        while s in seen:
            s, k = f"{base}_{k}", k + 1
        seen.add(s)
        out.append(s)
    return out


# =============================================================================
# Raw layer
# =============================================================================

def ensure_db(db: str) -> None:
    if db in PROTECTED:
        sys.exit(f"'{db}' is reserved; choose another dataset name")
    exists = ch(f"SELECT count() FROM system.databases WHERE name = '{db}'").strip() == "1"
    if exists and ch(f"EXISTS TABLE {db}._about").strip() != "1":
        sys.exit(f"database {db} already exists and was not made by auto_load; choose another name")
    ch(f"CREATE DATABASE IF NOT EXISTS {db}")
    ch(f"CREATE TABLE IF NOT EXISTS {db}._about (note String) ENGINE = Log")
    ch(f"CREATE TABLE IF NOT EXISTS {db}.raw (_source LowCardinality(String), _row UInt32, "
       f"_loaded_at DateTime DEFAULT now()) ENGINE = MergeTree PARTITION BY _source ORDER BY _row")


def raw_columns(db: str) -> list[str]:
    names = ch(f"SELECT name FROM system.columns WHERE database = '{db}' AND table = 'raw' "
               f"ORDER BY position FORMAT TSV").split()
    return [c for c in names if c not in RESERVED]


def stage(db: str, source: str, header: list[str], rows, replace_all: bool = False) -> int:
    source = source.replace("'", "")
    cols = sanitize(header)
    for c in cols:
        ch(f"ALTER TABLE {db}.raw ADD COLUMN IF NOT EXISTS `{c}` String DEFAULT ''")
    if replace_all:
        ch(f"TRUNCATE TABLE {db}.raw")
    else:
        ch(f"ALTER TABLE {db}.raw DROP PARTITION '{source}'")
    names = ["_source", "_row"] + cols
    batch, n = [], 0
    for r in rows:
        if not any(r):
            continue
        n += 1
        r = list(r) + [""] * (len(cols) - len(r))
        batch.append([source, n] + r[:len(cols)])
        if len(batch) == CHUNK:
            ax_load.insert(f"{db}.raw", names, batch)
            batch = []
    if batch:
        ax_load.insert(f"{db}.raw", names, batch)
    return n


# =============================================================================
# Deciding the model from the data
# =============================================================================

def profile(db: str, cols: list[str]) -> dict:
    pairs = ", ".join(f"('{c}', `{c}`)" for c in cols)
    q = f"""
    SELECT name, count() AS n, countIf(v != '') AS filled, uniqExactIf(v, v != '') AS distinct,
           countIf(v != '' AND toFloat64OrNull(v) IS NOT NULL AND isFinite(toFloat64OrZero(v))) AS numeric,
           countIf(match(v, '^-?\\\\d+$')) AS intlike,
           countIf(toFloat64OrZero(v) < 0) AS negative,
           max(abs(toFloat64OrZero(v))) AS maxabs,
           minIf(toFloat64OrZero(v), v != '') AS minv,
           countIf(match(v, '{DATE_RE}')) AS datelike,
           countIf(match(v, '^(1[3-9]|2\\\\d|3[01])[/.-]\\\\d{{1,2}}[/.-]')) AS dayfirst,
           countIf(match(v, '^\\\\d{{1,2}}[/.-](1[3-9]|2\\\\d|3[01])[/.-]')) AS monthfirst,
           countIf(match(v, '[eE][+-]\\\\d+$')) AS scientific,
           avgIf(length(v), v != '') AS avglen,
           countIf(position(v, ' ') > 0) AS spaced,
           groupUniqArrayIf(5)(v, v != '') AS sample
    FROM (SELECT arrayJoin([{pairs}]) AS p, p.1 AS name, p.2 AS v FROM {db}.raw)
    GROUP BY name FORMAT JSON"""
    return {r["name"]: r for r in json.loads(ch(q))["data"]}


def classify(name: str, s: dict) -> tuple[str, str, dict]:
    """(role, reason, extra). Roles: drop, date, measure, attribute, entity, document, text."""
    up, n, filled, distinct = name.upper(), int(s["n"]), int(s["filled"]), int(s["distinct"])
    few = max(50, min(LOWCARD, int(n * LOWCARD_SHARE)))
    if filled == 0:
        return "drop", "always empty", {}
    if distinct == 1 and filled == n:
        why = "the same value on every row"
        if DATEISH.search(up):
            why += f" ('{s['sample'][0]}') -- a date column with no usable dates (Excel formatting?)"
        return "drop", why, {}
    numeric, dates = int(s["numeric"]) / filled, int(s["datelike"]) / filled
    if numeric >= 0.98:
        lo, hi = float(s["minv"]), float(s["maxabs"])
        if 5.9e17 <= lo and hi <= 6.7e17:
            return "date", ".NET ticks (100ns since year 1)", {"kind": "ticks",
                                                             "rounded": int(s["scientific"]) > 0}
        if DATEISH.search(up) and 20000 <= lo and hi <= 80000:
            return "date", "Excel serial day numbers", {"kind": "excel"}
        if ID_SUFFIX.search(up):
            return _key_role(up, distinct, filled, few, "number, but named like a code or id")
        if MEASURE.search(up):
            return "measure", "number, named like an amount or quantity", {}
        # whole numbers with few values are codes -- also when a few letter codes
        # sit among them ('16', '12', 'BE', 'NS'): judged on the numeric values only
        if int(s["intlike"]) >= 0.99 * int(s["numeric"]) and (distinct <= 50 or hi >= 1e9):
            return _key_role(up, distinct, filled, few, "whole numbers with few values or id-sized")
        return "measure", "number with many different values (decided from values only)", {"stats_only": True}
    if dates >= 0.95:
        kind = "us" if int(s["monthfirst"]) > 0 and int(s["dayfirst"]) == 0 else "iso"
        return "date", "text dates" + (" (month/day/year)" if kind == "us" else ""), {"kind": kind}
    if distinct <= few:
        return "attribute", f"text with {distinct} different values", {}
    if float(s["avglen"]) > 40 or int(s["spaced"]) / filled > 0.5:
        return "text", "long free text", {}
    return _key_role(up, distinct, filled, few, "text id")


def _key_role(up: str, distinct: int, filled: int, few: int, why: str):
    if distinct <= few:
        return "attribute", f"{why}, {distinct} values", {}
    if DOCUMENT.search(up) or distinct / filled > 0.5:
        return "document", f"{why}: a document/line number, {distinct:,} values", {}
    return "entity", f"{why}: an entity (customer, meter, item...), {distinct:,} values", {}


def fd(db: str, a: str, bs: list[str]) -> set[str]:
    """Which b's are decided by a ('each a has one b'), over rows where a is filled.
    b must also be filled on at least half of those rows: a column that is
    nearly always empty looks 'decided' by anything, and belongs to nothing."""
    if not bs:
        return set()
    parts = ", ".join(f"uniqExact((`{a}`, `{b}`)), countIf(`{b}` != '')" for b in bs)
    vals = ch(f"SELECT uniqExact(`{a}`), count(), {parts} FROM {db}.raw WHERE `{a}` != '' FORMAT TSV").split()
    base, rows = int(vals[0]), int(vals[1])
    pairs = zip(bs, vals[2::2], vals[3::2])
    return {b for b, v, filled in pairs
            if int(v) - base <= FD_TOLERANCE * base and int(filled) >= 0.5 * rows}


def decide(db: str, name: str, source: dict) -> dict:
    cols = raw_columns(db)
    stats = profile(db, cols)
    nrows = int(ch(f"SELECT count() FROM {db}.raw").strip())
    columns, check = {}, []
    for c in cols:
        role, why, extra = classify(c, stats[c])
        columns[c] = {"role": role, "why": why, **extra}
    by = lambda role: [c for c in cols if columns[c]["role"] == role]
    fill = {c: int(stats[c]["filled"]) for c in cols}
    distinct = {c: int(stats[c]["distinct"]) for c in cols}

    # --- the date that drives time --------------------------------------------
    dates = by("date")
    primary = None
    if dates:
        primary = max(dates, key=lambda c: (fill[c], bool(PREFERRED_DATE.search(c.upper()))))
        span = ch(f"SELECT dateDiff('day', min(d), max(d)) FROM (SELECT {date_expr(primary, columns[primary])} AS d "
                  f"FROM {db}.raw) WHERE d > '1970-01-01'").strip()
        grain = "month" if span and int(span) > 62 else "day"
        if columns[primary].get("rounded"):
            check.append(f"Dates come from `{primary}`, rounded by Excel to ~1 day: a few rows near a "
                         f"month end may land in the neighbouring month.")
    else:
        grain = None
        check.append("No usable date column: aggregates have no time axis. Ask for an export with dates.")
    date_fallback = None
    if primary:
        d = date_expr(primary, columns[primary])
        per = f"toStartOfMonth({d})" if grain == "month" else d
        undated = ch(f"SELECT count(), round(sum(toFloat64OrZero(`{by('measure')[0]}`)), 2) FROM {db}.raw "
                     f"WHERE {d} = '1970-01-01' FORMAT TSV" if by("measure") else
                     f"SELECT count(), 0 FROM {db}.raw WHERE {d} = '1970-01-01' FORMAT TSV").split()
        if int(undated[0]):
            # a column whose every value sits (99%+) in one period -- e.g. a batch or
            # run id -- tells the period of the undated rows too
            best = 0.0
            for c in by("attribute"):
                if not (2 <= distinct[c] <= 200 and fill[c] >= 0.99 * nrows):
                    continue
                share = float(ch(f"SELECT sum(top) / sum(total) FROM (SELECT max(n) AS top, sum(n) AS total FROM "
                                 f"(SELECT `{c}` AS k, {per} AS p, count() AS n FROM {db}.raw WHERE {d} > '1970-01-01' "
                                 f"GROUP BY k, p) GROUP BY k)").strip() or 0)
                if share >= 0.99 and share > best:
                    best, date_fallback = share, c
            what = f"{int(undated[0]):,} rows" + (f" (`{by('measure')[0]}` {float(undated[1]):,.2f})" if by("measure") else "")
            if date_fallback:
                check.append(f"{what} have no date: each takes the {grain} most common for its `{date_fallback}` "
                             f"({best:.1%} of dated rows follow that rule).")
            else:
                check.append(f"{what} have no date and nothing predicts one: they sit in period 1970-01-01.")

    # --- dimensions: code/description pairs and hierarchies, found in the data --
    attrs = sorted(by("attribute"),
                   key=lambda c: (-distinct[c], not ID_SUFFIX.search(c.upper()), float(stats[c]["avglen"])))
    roots, taken = [], set()
    for a in attrs:
        if a in taken:
            continue
        # a sparse column must not swallow a fuller one (a water-only node cannot
        # own the region of every row)
        cands = [b for b in attrs if b != a and b not in taken
                 and distinct[b] <= distinct[a] and fill[b] <= fill[a]]
        deps = sorted(fd(db, a, cands), key=attrs.index)
        roots.append({"key": a, "attributes": deps})
        taken |= {a, *deps}
    entities = []
    for e in sorted(by("entity"), key=lambda c: distinct[c]):
        cands = [r["key"] for r in roots if fill[r["key"]] <= fill[e]]
        entities.append({"key": e, "attributes": sorted(fd(db, e, cands))})

    # --- measures -------------------------------------------------------------
    measures = [{"name": m, "additive": not NON_ADDITIVE.search(m.upper())} for m in by("measure")]
    for m in measures:
        st = stats[m["name"]]
        if int(st["negative"]):
            check.append(f"`{m['name']}` has {int(st['negative']):,} negative values (credits, reversals?): "
                         f"totals are net of them.")
        if not m["additive"]:
            check.append(f"`{m['name']}` looks like a price/rate: aggregates AVERAGE it instead of summing.")
        if columns[m["name"]].get("stats_only"):
            check.append(f"`{m['name']}` was made a measure from its values alone -- if it is really a code "
                         f"or id, set its role to \"attribute\" in models/{name}.json and --rebuild.")
    # quantities in different units: price per unit differs hugely between groups
    money = [m["name"] for m in measures if MONEY.search(m["name"].upper())]
    for q in [m["name"] for m in measures if QTY.search(m["name"].upper())]:
        for r in [r["key"] for r in roots if 2 <= distinct[r["key"]] <= 20]:
            if not money:
                break
            rates = [float(x) for x in ch(
                f"SELECT sum(toFloat64OrZero(`{money[0]}`)) / sum(toFloat64OrZero(`{q}`)) FROM {db}.raw "
                f"GROUP BY `{r}` HAVING sum(toFloat64OrZero(`{q}`)) > 0 FORMAT TSV").split()]
            rates = [x for x in rates if x > 0]
            if len(rates) >= 2 and max(rates) / min(rates) > 3:
                next(x for x in measures if x["name"] == q).setdefault("split_by", []).append(r)
        split = next(x for x in measures if x["name"] == q).get("split_by", [])
        if split:
            # the most basic splitter (fewest values, always filled) is taken as the
            # unit: e.g. utility type (kWh vs m3) rather than a charge type
            split.sort(key=lambda r: (distinct[r], -fill[r]))
            check.append(f"`{q}` has very different prices per unit across {', '.join(f'`{r}`' for r in split)} "
                         f"-- probably different units (kWh vs m3?). Never add `{q}` up across "
                         f"`{split[0]}`: every aggregate keeps it. If the unit is really set by another "
                         f"column, move that one first in \"split_by\" in models/{name}.json and --rebuild.")
    for r in roots:
        if not r["attributes"] and int(stats[r["key"]]["intlike"]) == fill[r["key"]] and distinct[r["key"]] > 1:
            check.append(f"`{r['key']}` is a code with no description that matches it one-to-one "
                         f"(values {', '.join(stats[r['key']]['sample'][:5])}): ask what the codes mean.")
    dupes = int(ch(f"SELECT count() - uniqExact(cityHash64({', '.join(f'`{c}`' for c in cols)})) "
                   f"FROM {db}.raw").strip())
    if dupes:
        check.append(f"{dupes:,} rows are exact copies of another row. Kept (they can be real separate "
                     f"lines), but a line id in the export would make them tell-apart-able.")
    for c in by("drop"):
        if "Excel" in columns[c]["why"]:
            check.append(f"`{c}` is unusable: {columns[c]['why']}.")

    return {"name": name, "created": datetime.now().isoformat(timespec="seconds"), "rows_seen": nrows,
            "source": source, "columns": columns, "date": primary, "grain": grain,
            "date_fallback": date_fallback,
            "measures": measures, "dimensions": roots, "entities": entities,
            "aggregates": None, "check": check}


# =============================================================================
# Building fact, dimensions and aggregates from the model
# =============================================================================

def date_expr(c: str, meta: dict) -> str:
    v, kind = f"`{c}`", meta.get("kind")
    if kind == "ticks":
        e = f"if(toFloat64OrZero({v}) > 0, toDate(toDateTime(toInt64(toFloat64OrZero({v}) / 1e7 - 62135596800))), NULL)"
    elif kind == "excel":
        e = f"if(toFloat64OrZero({v}) > 0, toDate('1899-12-30') + toInt32(toFloat64OrZero({v})), NULL)"
    else:
        fn = "parseDateTimeBestEffortUSOrNull" if kind == "us" else "parseDateTimeBestEffortOrNull"
        # the pattern check first: best-effort parsing turns '00:00.0' into today
        e = f"if(match({v}, '{DATE_RE}'), toDate({fn}({v})), NULL)"
    return f"ifNull({e}, toDate('1970-01-01'))"


def period_expr(model: dict, db: str | None = None) -> str | None:
    if not model["date"]:
        return None
    d = date_expr(model["date"], model["columns"][model["date"]])
    per = f"toStartOfMonth({d})" if model["grain"] == "month" else d
    fb = model.get("date_fallback")
    if not (fb and db):
        return per
    # undated rows: the period most common for their value of the fallback column,
    # recomputed every build so new batches are covered
    pairs = [l.split("\t") for l in ch(
        f"SELECT `{fb}`, toString(topK(1)({per})[1]) FROM {db}.raw WHERE {d} > '1970-01-01' "
        f"GROUP BY `{fb}` FORMAT TSV").splitlines() if l]
    if not pairs:
        return per
    keys = ", ".join("'" + k.replace("\\", "\\\\").replace("'", "\\'") + "'" for k, _ in pairs)
    vals = ", ".join(f"toDate('{v}')" for _, v in pairs)
    return f"if({d} > '1970-01-01', {per}, transform(`{fb}`, [{keys}], [{vals}], toDate('1970-01-01')))"


def build_fact(db: str, model: dict) -> None:
    present = set(raw_columns(db))
    sel = ["_source", "_row"]
    p = period_expr(model, db)
    if p:
        sel.append(f"{p} AS period")
    for c, meta in model["columns"].items():
        if c not in present or meta["role"] == "drop":
            continue
        if meta["role"] == "date":
            sel.append(f"{date_expr(c, meta)} AS `{c}`")
        elif meta["role"] == "measure":
            # float -> text -> Decimal: a direct float -> Decimal cast truncates 1070.37 to 1070.3699
            sel.append(f"toDecimal64OrZero(toString(round(toFloat64OrZero(`{c}`), 4)), 4) AS `{c}`")
        elif meta["role"] == "attribute":
            sel.append(f"toLowCardinality(`{c}`) AS `{c}`")
        else:
            sel.append(f"`{c}`")
    new = present - set(model["columns"])
    if new:
        print(f"  note: columns not in the model are kept in raw only: {', '.join(sorted(new))} "
              f"(--remodel to include them)")
    order = "(period, _source, _row)" if p else "(_source, _row)"
    ch(f"CREATE OR REPLACE TABLE {db}.fact ENGINE = MergeTree ORDER BY {order} AS "
       f"SELECT {', '.join(sel)} FROM {db}.raw SETTINGS prefer_column_name_to_alias = 1")


def uniq(db: str, cols: list[str]) -> int:
    return int(ch(f"SELECT uniqExact(tuple({', '.join(cols)})) FROM {db}.fact").strip())


def choose_aggregates(db: str, model: dict) -> dict:
    """Pick grains from the data: the main aggregate keeps every dimension that
    still collapses rows well; one per entity (customer, meter...) as well."""
    nrows = int(ch(f"SELECT count() FROM {db}.fact").strip())
    budget = max(1000, min(200_000, nrows // 10))
    base = ["period"] if model["date"] else []
    must = sorted({r for m in model["measures"] for r in m.get("split_by", [])})       # main aggregate
    unit = sorted({m["split_by"][0] for m in model["measures"] if m.get("split_by")})  # everywhere
    keys = sorted((r["key"] for r in model["dimensions"] if r["key"] not in must),
                  key=lambda k: uniq(db, [f"`{k}`"]))
    keys_all = keys + [k for k in must if k not in unit]
    main, left_out = base + [f"`{k}`" for k in must], []
    for k in keys:
        if uniq(db, main + [f"`{k}`"]) <= budget:
            main.append(f"`{k}`")
        else:
            left_out.append(k)
    per_entity = []
    for e in model["entities"][:2]:  # the two broadest entities, e.g. customer and meter point
        grain = base + [f"`{e['key']}`"] + [f"`{k}`" for k in unit]
        start = uniq(db, grain)
        for k in keys_all:
            if f"`{k}`" not in grain and uniq(db, grain + [f"`{k}`"]) <= 1.5 * start:
                grain.append(f"`{k}`")
        per_entity.append({"entity": e["key"], "grain": [g.strip("`") for g in grain]})
    return {"main": [g.strip("`") for g in main], "left_out": left_out, "per_entity": per_entity}


def agg_select(model: dict, grain: list[str], counts: list[str]) -> str:
    cols = [f"`{g}`" for g in grain]
    for m in model["measures"]:
        if m["additive"]:
            cols.append(f"toDecimal64(sum(`{m['name']}`), 4) AS `{m['name']}`")
        else:
            cols.append(f"round(avg(toFloat64(`{m['name']}`)), 4) AS `{m['name']}_avg`")
    cols.append("count() AS row_count")
    cols += [f"uniqExact(`{e}`) AS `{e}_count`" for e in counts]
    return ", ".join(cols)


def build_rest(db: str, model: dict) -> list[str]:
    made = []
    for r in model["dimensions"]:
        if not r["attributes"]:
            continue
        deps = ", ".join(f"topK(1)(`{a}`)[1] AS `{a}`" for a in r["attributes"])
        ch(f"CREATE OR REPLACE TABLE {db}.dim_{r['key']} ENGINE = MergeTree ORDER BY `{r['key']}` AS "
           f"SELECT `{r['key']}`, {deps}, count() AS row_count FROM {db}.fact GROUP BY `{r['key']}`")
        made.append(f"dim_{r['key']}")
    for e in model["entities"]:
        extra = "".join(f", argMax(`{a}`, (period, _row)) AS `{a}`" if model["date"] else f", any(`{a}`) AS `{a}`"
                        for a in e["attributes"])
        span = ", min(period) AS first_period, max(period) AS last_period" if model["date"] else ""
        ch(f"CREATE OR REPLACE TABLE {db}.dim_{e['key']} ENGINE = MergeTree ORDER BY `{e['key']}` AS "
           f"SELECT `{e['key']}`, count() AS row_count{span}{extra} FROM {db}.fact GROUP BY `{e['key']}`")
        made.append(f"dim_{e['key']}")
    lo, hi = (ch(f"SELECT minIf(period, period > '1970-01-01'), max(period) FROM {db}.fact FORMAT TSV").split()
              if model["date"] else ("1970-01-01", "1970-01-01"))
    if lo != "1970-01-01":  # a calendar from the first to the last month in the data
        days = (date.fromisoformat(hi) - date.fromisoformat(lo)).days + 32
        ch(f"""CREATE OR REPLACE TABLE {db}.dim_date ENGINE = MergeTree ORDER BY date_key AS
               SELECT toDate('{lo}') + number AS date_key, toYear(date_key) AS year,
                      toQuarter(date_key) AS quarter, toMonth(date_key) AS month,
                      formatDateTime(date_key, '%b') AS month_name,
                      formatDateTime(date_key, '%Y-%m') AS year_month, toStartOfMonth(date_key) AS month_start
               FROM numbers({days})""")
        made.append("dim_date")

    aggs = model["aggregates"]
    counts = [e["entity"] for e in aggs["per_entity"]]
    suffix = {"month": "monthly", "day": "daily"}.get(model["grain"], "")
    main = f"agg_{suffix}" if model["date"] else "agg_summary"
    ch(f"CREATE OR REPLACE TABLE {db}.{main} ENGINE = MergeTree ORDER BY ({', '.join(f'`{g}`' for g in aggs['main']) or 'tuple()'}) "
       f"AS SELECT {agg_select(model, aggs['main'], counts)} FROM {db}.fact "
       f"GROUP BY {', '.join(f'`{g}`' for g in aggs['main']) or 'tuple()'} SETTINGS prefer_column_name_to_alias = 1")
    made.append(main)
    for pe in aggs["per_entity"]:
        t = f"agg_{pe['entity']}_{suffix}" if model["date"] else f"agg_{pe['entity']}"
        g = ", ".join(f"`{x}`" for x in pe["grain"])
        ch(f"CREATE OR REPLACE TABLE {db}.{t} ENGINE = MergeTree ORDER BY ({g}) "
           f"AS SELECT {agg_select(model, pe['grain'], [])} FROM {db}.fact GROUP BY {g} "
           f"SETTINGS prefer_column_name_to_alias = 1")
        made.append(t)
    ch(f"CREATE OR REPLACE VIEW {db}.v_sources AS SELECT _source, count() AS rows, max(_loaded_at) AS loaded_at "
       f"FROM {db}.raw GROUP BY _source")
    return made


def report(model: dict, made: list[str]) -> str:
    cols = model["columns"]
    by = lambda role: [c for c in cols if cols[c]["role"] == role]
    L = [f"# {model['name']}: what was decided", "",
         f"Model made {model['created']} from {model['rows_seen']:,} rows. Every later load reuses it.", ""]
    if model["date"]:
        L += [f"**Time:** `{model['date']}` ({cols[model['date']]['why']}), aggregated by **{model['grain']}**.", ""]
    L += ["**Measures (summed):** " + (", ".join(f"`{m['name']}`" for m in model["measures"] if m["additive"]) or "none"), ""]
    avg = [m["name"] for m in model["measures"] if not m["additive"]]
    if avg:
        L += ["**Measures (averaged):** " + ", ".join(f"`{a}`" for a in avg), ""]
    L += ["**Dimensions** (key → attributes it decides):", ""]
    for r in model["dimensions"]:
        L.append(f"- `{r['key']}`" + (" → " + ", ".join(f"`{a}`" for a in r["attributes"]) if r["attributes"] else ""))
    L += ["", "**Entities** (many values, own dimension table + aggregate):", ""]
    L += [f"- `{e['key']}`" + (" → " + ", ".join(f"`{a}`" for a in e["attributes"]) if e["attributes"] else "")
          for e in model["entities"]] or ["- none"]
    for role, title in (("document", "Document numbers (kept on the fact for drill-down)"),
                        ("text", "Free text (fact only)"), ("drop", "Dropped")):
        if by(role):
            L += ["", f"**{title}:**", ""] + [f"- `{c}`: {cols[c]['why']}" for c in by(role)]
    aggs = model["aggregates"]
    L += ["", "**Tables:** `raw`, `fact`, " + ", ".join(f"`{t}`" for t in made), "",
          "Main aggregate grain: " + ", ".join(f"`{g}`" for g in aggs["main"])]
    if aggs["left_out"]:
        L.append("Too detailed for it (use `fact`): " + ", ".join(f"`{k}`" for k in aggs["left_out"]))
    L += ["", "## Check", ""] + ([f"- {c}" for c in model["check"]] or ["- nothing flagged"])
    L += ["", "Distinct counts (`*_count`) in aggregates are per row: never sum them across rows."]
    return "\n".join(L) + "\n"


def build(name: str, source: dict | None = None, remodel: bool = False) -> None:
    db = name
    path = MODELS / f"{name}.json"
    t0 = time.time()
    if remodel or not path.exists():
        print("  deciding the model from the data...")
        model = decide(db, name, source or {})
        build_fact(db, model)
        model["aggregates"] = choose_aggregates(db, model)
        MODELS.mkdir(exist_ok=True)
        path.write_text(json.dumps(model, indent=2))
    else:
        model = json.loads(path.read_text())
        build_fact(db, model)
    made = build_rest(db, model)
    (MODELS / f"{name}.md").write_text(report(model, made))
    rows = ch(f"SELECT count() FROM {db}.fact").strip()
    print(f"  {db}: fact {int(rows):,} rows, {', '.join(made)}  ({time.time() - t0:.0f}s)")
    print(f"  report: models/{name}.md" + (f"  ({len(model['check'])} things to check)" if model["check"] else ""))


# =============================================================================
# Entry points
# =============================================================================

def load_file(name: str, path: Path, remodel: bool = False) -> None:
    ensure_db(name)
    for source, header, rows in read_any(path):
        t0 = time.time()
        n = stage(name, source, header, rows)
        print(f"{name} <- {source}: {n:,} rows ({time.time() - t0:.0f}s)")
    build(name, {"kind": "file", "last_file": path.name}, remodel)


def load_sql(name: str, sql: str, db: str | None, remodel: bool = False) -> None:
    import os
    if db:
        os.environ["MSSQL_DB"] = db
    ensure_db(name)
    t0 = time.time()
    conn = ax_load.connect()
    cur = conn.cursor()
    cur.execute(sql)
    header = [d[0] for d in cur.description]

    def rows():
        while batch := cur.fetchmany(CHUNK):
            yield from ([cell(v) for v in r] for r in batch)

    n = stage(name, "sql", header, rows(), replace_all=True)
    conn.close()
    print(f"{name} <- SQL Server: {n:,} rows ({time.time() - t0:.0f}s)")
    build(name, {"kind": "sql", "sql": sql, "db": db or os.environ.get("MSSQL_DB")}, remodel)


def dataset_name(inbox: Path, f: Path) -> str:
    if f.parent != inbox:
        return sanitize([f.parent.name])[0]
    return sanitize([re.sub(r"\d+", "", f.stem)])[0]


def run_inbox(inbox: Path) -> None:
    inbox = inbox.expanduser().resolve()
    files = [f for f in sorted(inbox.rglob("*"))
             if f.is_file() and f.suffix.lower() in (".csv", ".tsv", ".txt", ".xlsx", ".xlsm", ".zip")
             and not {"done", "failed"} & set(f.relative_to(inbox).parts[:1])]
    if not files:
        return
    for f in files:
        name = dataset_name(inbox, f)
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        try:
            load_file(name, f)
            dest = inbox / "done" / name
            dest.mkdir(parents=True, exist_ok=True)
            shutil.move(str(f), dest / f"{stamp}_{f.name}")
        except (Exception, SystemExit) as e:
            dest = inbox / "failed"
            dest.mkdir(parents=True, exist_ok=True)
            shutil.move(str(f), dest / f"{stamp}_{f.name}")
            (dest / f"{stamp}_{f.name}.error.txt").write_text(str(e))
            print(f"FAILED {f.name}: {str(e).splitlines()[0] if str(e) else e!r}")


def main() -> None:
    ap = argparse.ArgumentParser(description="Load raw data into ClickHouse and aggregate it automatically.")
    ap.add_argument("name", nargs="?", help="dataset name = ClickHouse database, e.g. billing")
    src = ap.add_mutually_exclusive_group()
    src.add_argument("--file", type=Path)
    src.add_argument("--sql")
    src.add_argument("--sql-file", type=Path)
    src.add_argument("--inbox", type=Path)
    src.add_argument("--rebuild", action="store_true")
    ap.add_argument("--db", help="SQL Server database for --sql (default MSSQL_DB from .env)")
    ap.add_argument("--remodel", action="store_true", help="decide the model again from the data")
    a = ap.parse_args()

    lock = open(HERE / ".auto_load.lock", "w")
    fcntl.flock(lock, fcntl.LOCK_EX)  # cron and a manual run never collide
    if a.inbox:
        return run_inbox(a.inbox)
    if not a.name:
        ap.error("give a dataset name, e.g.  python3 auto_load.py billing --file report.csv")
    name = sanitize([a.name])[0]
    if a.file:
        load_file(name, a.file, a.remodel)
    elif a.sql or a.sql_file:
        load_sql(name, a.sql or a.sql_file.read_text(), a.db, a.remodel)
    elif a.rebuild or a.remodel:
        ensure_db(name)
        build(name, remodel=a.remodel)
    else:
        ap.error("say where the data comes from: --file, --sql, --sql-file or --inbox")


if __name__ == "__main__":
    main()
